import streamlit as st
import openpyxl
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles.borders import Border, Side
from openpyxl.styles import PatternFill, Font
st.set_page_config(layout="wide", page_title="Product Change Report")
st.write("## Product change report for GL")
st.write("Upload current PAS.xlsx, GP.xlsx, GPWC.xlsx, JSTR.xlsx, KMET.xlsx, BOSTCO.xlsx, synonyms.xlsx and previous product change report files to generate a Product Change Report.")
data = st.file_uploader('Multiple Excel files', type=["xlsx"], accept_multiple_files=True)
Current_tank_list = ["PAS.xlsx","GP.xlsx", "GPWC.xlsx", "JSTR.xlsx", "KMET.xlsx", "BOSTCO.xlsx"]
old_rpt = st.sidebar.text_input("Previous Product Change Report Date (mm-dd-yyyy):")
new_rpt = st.sidebar.text_input("Current Report Date (mm-dd-yyyy):")
def process_product_change():
    # Synonyms file
    synonyms=pd.read_excel(r'synonyms.xlsx',sheet_name="Chemicals 2024")
    synonyms['SYNONYM']=synonyms['SYNONYM'].str.upper()
    # Combined report workbook file
    ProductChange = 'Product Change '+ new_rpt + '.xlsx'
    Previous_week = 'Product Change '+ old_rpt + '.xlsx'
    # report workbook before being saved
    Rpt = Workbook()
    # read all tank inventory excel files
    for Inv_file in Current_tank_list:
        # open each product inventory workbook
        Inv = openpyxl.load_workbook(Inv_file)
        Inv_sheet = Inv.active
        # create a new report worksheet with the name of the product inventory workbook file name    
        Rpt_sheet = Rpt.create_sheet(Inv_file.replace('.xlsx',''))
        for i in range (1, Inv_sheet.max_row + 1 ):
            for j in range (1,Inv_sheet.max_column + 1):
                # read each cell value from product inventory worksheet            
                k = Inv_sheet.cell(row = i, column = j).value
                # write the copied value into the new report worksheet
                Rpt_sheet.cell(row = i, column = j).value = k
                # add border to each cell           
                Rpt_sheet.cell(row = i, column = j).border = Border(
                                                                        left=Side(style='thin'), 
                                                                        right=Side(style='thin'), 
                                                                        top=Side(style='thin'), 
                                                                        bottom=Side(style='thin')
                                                                        )       
        # adjust column width   
        Rpt_sheet.column_dimensions['A'].width = 15
        Rpt_sheet.column_dimensions['B'].width = 20
        Rpt_sheet.column_dimensions['C'].width = 20  
        Rpt_sheet.column_dimensions['D'].width = 15  
        Rpt_sheet.column_dimensions['E'].width = 15
    # delete the extra blank tab
    del Rpt['Sheet']
    # delete extra rows and columns from the original tank inventory files
    for m in range (0,6):
        Rpt.worksheets[m].delete_rows(1,4)
        Rpt.worksheets[m].delete_cols(2)
        Rpt.worksheets[m].delete_cols(4)
        Rpt.worksheets[m].delete_cols(5)
        Rpt.worksheets[m].delete_cols(6,2)
    # fill first row with yellow color and apply bold font to tank inventory tabs for each site   
        for n in ['A1','B1','C1','D1','E1']:
            Rpt.worksheets[m][n].fill = PatternFill('solid', start_color='00FFFFCC')
            Rpt.worksheets[m][n].font = Font(bold=True)
    Rpt.save(ProductChange)
    Rpt.close()
    #import current week data into pandas
    Current_tank_inventory = pd.read_excel(ProductChange,sheet_name=['BOSTCO','GP','GPWC','JSTR','KMET','PAS'])
    bostco = Current_tank_inventory['BOSTCO']
    gp = Current_tank_inventory['GP']
    gpwc = Current_tank_inventory['GPWC']
    jstr = Current_tank_inventory['JSTR']
    kmet = Current_tank_inventory['KMET']
    pas = Current_tank_inventory['PAS']
    #add a new location column
    gp['Location']='GP'
    gpwc['Location']='GPWC'
    jstr['Location']='JSTR'
    kmet['Location']='KMET'
    pas['Location']='PAS'
    bostco['Location']='BOSTCO'
    #merge all of the frames into gulf liquids region
    frames=[gp,gpwc,jstr,kmet,pas,bostco]
    nonewsyn=pd.concat(frames)

    # Check if all current week products are in synonym list
    synonym_verification = set(nonewsyn['PRODUCT']) - set(synonyms['SYNONYM'])
    if len(synonym_verification) == 0:
        st.write('Congratulations! No new synonyms this week!')
    for item in synonym_verification:
        st.write(f"'{item}' is not in the synonyms, check the result manually!!!")
    #Ensure datatypes remain the same during merge, remove paces
    synonyms=synonyms.rename({'TERMINAL_NAME':'Location', 'SYNONYM':'PRODUCT'}, axis=1)
    synonyms['Location']=synonyms['Location'].astype(str)
    synonyms['Location']=synonyms['Location'].str.strip()

    newsyn=pd.merge(nonewsyn, synonyms, on=['PRODUCT', 'Location'], how='left')

    glr=newsyn[['Tank Name', 'Customer', 'PRODUCT', 'Current Book', 'UOM', 'Location', 'Service', 'OLD']]

    #certain tanks are in the wrong place correct these
    #tank 130-1 is labelled as JSTR, it should be labeled as PAS
    #80-101 and 80-102 need to be deleted from JSTR, these are gulf central tanks
    #~ means the function will do the opposite
    pastank=glr['Tank Name']=='130-1'
    glr.loc[pastank,'Location']='PAS'
    glr2=glr[(glr['Tank Name']=='80-101') | (glr['Tank Name']=='80-102')]
    glr=glr[~glr.isin(glr2)].dropna()
    glr=glr.rename(columns={"PRODUCT": "Product"})
    glr['Date']= new_rpt
    glr['Tank Name']=glr['Tank Name'].astype('string')

    #import previous week data into pandas
    Previous_tank_inventory = pd.read_excel(Previous_week,sheet_name=['BOSTCO','GP','GPWC','JSTR','KMET','PAS'])
    oldbostco = Previous_tank_inventory['BOSTCO']
    oldgp = Previous_tank_inventory['GP']
    oldgpwc = Previous_tank_inventory['GPWC']
    oldjstr = Previous_tank_inventory['JSTR']
    oldkmet = Previous_tank_inventory['KMET']
    oldpas = Previous_tank_inventory['PAS']
    oldgp['Location']='GP'
    oldgpwc['Location']='GPWC'
    oldjstr['Location']='JSTR'
    oldkmet['Location']='KMET'
    oldpas['Location']='PAS'
    oldbostco['Location']='BOSTCO'
    framesold=[oldgp,oldgpwc,oldjstr,oldkmet,oldpas,oldbostco]

    #merge all of the frames into gulf liquids region
    #add old date
    nooldsyn=pd.concat(framesold)

    oldsyn=pd.merge(nooldsyn, synonyms, on=['Location', 'PRODUCT'], how='left')
    oldglr=oldsyn[['Tank Name', 'Customer', 'PRODUCT', 'Current Book', 'UOM', 'Location', 'Service', 'OLD']]

    pastankold=oldglr['Tank Name']=='130-1'
    oldglr.loc[pastankold,'Location']='PAS'
    oldglr2=oldglr[(oldglr['Tank Name']=='80-101') | (oldglr['Tank Name']=='80-102')]
    oldglr=oldglr[~oldglr.isin(oldglr2)].dropna()
    oldglr=oldglr.rename(columns={"PRODUCT": "Product"})
    oldglr['Date']= old_rpt
    oldglr=oldglr.replace('-', np.NaN)
    oldglr['Tank Name']=oldglr['Tank Name'].astype('string')

    merged=pd.merge(glr,oldglr, on = ['Tank Name', 'Location', 'Product'], how='outer', indicator=True)
    changesonly=merged[(merged['_merge']=='right_only') | (merged['_merge']=='left_only')]
    new=merged[(merged['_merge']=='right_only')]
    old=merged[(merged['_merge']=='left_only')]
    merged2=pd.merge(new, old, on = ['Tank Name','Location'], how='outer')
    intermediate=merged2[(merged2['_merge_y']=='left_only')]
    intermediate3=intermediate[['Location', 'Tank Name', 'Product_x', 'Service_y_x','OLD_y_x','Product_y','Service_x_y', 'OLD_x_y']].drop_duplicates()
    summary=intermediate3.rename(columns={"Product_x": old_rpt, "Service_y_x":"Old Service","OLD_y_x":"Previous OLD Status", "Product_y":new_rpt, "Service_x_y":"New Service", "OLD_x_y":"New OLD Status"})
    with pd.ExcelWriter(
        ProductChange,
        mode="a",
        engine="openpyxl",
        if_sheet_exists="replace",
    ) as writer:
        summary.to_excel(writer, sheet_name='Product Change '+new_rpt,index=False,header=True)
    ProductChange_wb=openpyxl.load_workbook(ProductChange)
    summary_ws = ProductChange_wb['Product Change '+ new_rpt]
    # t = summary_ws.cell(row = 3, column = 3).value
    # print(t)
    for a in range (1, summary_ws.max_row + 1 ):
            for b in range (1,summary_ws.max_column + 1):
                # add border to each cell           
                summary_ws.cell(row = a, column = b).border = Border(
                                                                        left=Side(style='thin'), 
                                                                        right=Side(style='thin'), 
                                                                        top=Side(style='thin'), 
                                                                        bottom=Side(style='thin')
                                                                        )       
    # adjust column width   
    summary_ws.column_dimensions['A'].width = 15
    summary_ws.column_dimensions['B'].width = 15
    summary_ws.column_dimensions['C'].width = 20
    summary_ws.column_dimensions['D'].width = 15
    summary_ws.column_dimensions['E'].width = 20
    summary_ws.column_dimensions['F'].width = 20
    summary_ws.column_dimensions['G'].width = 15
    summary_ws.column_dimensions['H'].width = 20
    #fill first row with yellow color and apply bold font to product change summary tab
    for x in ['A1','B1','C1','D1','E1','F1', 'G1', 'H1']:
        ProductChange_wb.worksheets[6][x].fill = PatternFill('solid', start_color='00FFFFCC')
        ProductChange_wb.worksheets[6][x].font = Font(bold=True)
    ProductChange_wb.save(ProductChange)
    ProductChange_wb.close()
    # output = BytesIO()
    # ProductChange_wb = xlsxwriter.Workbook(output, {'in_memory': True})
    # st.sidebar.download_button(
    # label="Download Product Change Report",
    # data=output.getvalue(),
    # file_name='Product Change '+ new_rpt+'.xlsx',
    # mime="application/vnd.ms-excel"
    # )
    # b64 = base64.b64encode(ProductChange).decode('UTF-8')
    # linko_final= f'<a href="data:file/xlsx base64,{b64}" download={ProductChange}>Download Product Change Report</a>'
    # st.markdown(linko_final, unsafe_allow_html=True)  
if st.sidebar.button("Run Product Change Report",type="primary"):
    process_product_change()
